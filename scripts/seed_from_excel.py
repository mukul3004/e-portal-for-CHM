import os
import sys
import uuid
import hashlib
from datetime import datetime, time, timedelta, date
from decimal import Decimal
import openpyxl

# Add project root to sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from backend.app.database import SessionLocal, engine, Base
from backend.app.models.user import User, RoleType
from backend.app.models.case import Case, CaseType, CaseStatus, Pleading, PleadingType, Exhibit
from backend.app.models.admissions import AdmissionsChecklist, AdmissionItem, StanceType, ItemResolutionStatus
from backend.app.models.schedule import CMHSession, SessionStatus
from backend.app.models.order import CaseManagementOrder, DiscoveryDeadline, DeadlineStatus

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
LOCAL_DATA_PATH = os.path.join(BASE_DIR, "data", "CMH_eportal_dataset.xlsx")
DOWNLOAD_PATH = "/Users/mukulbhamgare/Downloads/CMH_eportal_dataset.xlsx"
EXCEL_PATH = LOCAL_DATA_PATH if os.path.exists(LOCAL_DATA_PATH) else DOWNLOAD_PATH

CATEGORY_MAP = {
    "Commercial Disputes": CaseType.COMMERCIAL,
    "Cheque Bounce": CaseType.CHEQUE_BOUNCE,
    "Arbitration": CaseType.ARBITRATION,
    "Contract Breach": CaseType.CONTRACT,
    "Civil Suits": CaseType.CIVIL_RECOVERY
}

COURT_CNR_PREFIX = {
    "MP High Court, Indore Bench": "MPHC01",
    "Commercial Court, Indore": "MPIN02",
    "District Court, Indore": "MPIN01",
    "District Court, Bhopal": "MPBP01",
    "District Court, Gwalior": "MPGW01",
    "District Court, Ujjain": "MPUJ01",
}

def clean_name(val):
    if not val:
        return ""
    return str(val).replace("(Synthetic)", "").strip()

def get_averments_for_category(category, plaintiff, defendant, statute, claim_val):
    if category == "Cheque Bounce":
        return [
            f"The complainant {plaintiff} provided financial facilities/goods to {defendant} against legitimate commercial liability.",
            f"Cheque drawn on State Bank of India for INR {claim_val:,.2f} was signed and issued by authorized signatory of {defendant}.",
            f"Cheque was presented within validity and dishonoured with bank return memo stating 'Funds Insufficient'.",
            f"Statutory demand notice under Section 138(b) of Negotiable Instruments Act, 1881 was dispatched within 30 days.",
            f"Accused failed to make payment within 15 days of statutory notice receipt."
        ]
    elif category == "Arbitration":
        return [
            f"Parties entered into commercial agreement governed by arbitration clause under {statute}.",
            f"Dispute arose regarding final reconciliation and performance covenants.",
            f"Notice invoking arbitration under Section 21 was formally served on {defendant}.",
            f"Impugned arbitral award was delivered by the Arbitral Tribunal with reasoned findings.",
            f"The Section 34 petition has been presented within the statutory 90-day limitation window."
        ]
    elif category == "Contract Breach":
        return [
            f"Master Services and Supply Agreement was executed between {plaintiff} and {defendant}.",
            f"{plaintiff} fulfilled preliminary delivery conditions and supplied goods/services.",
            f"{defendant} failed to adhere to milestone payment obligations exceeding INR {claim_val:,.2f}.",
            f"Cure notice was issued providing 30 days remediation opportunity.",
            f"Liquidated damages claimed represent genuine agreed pre-estimate of losses."
        ]
    elif category == "Civil Suits":
        return [
            f"Agreement to Sell was executed regarding commercial property in Indore.",
            f"{plaintiff} deposited earnest money and has always been ready and willing to perform contract obligations.",
            f"Notice requiring execution of registered conveyance deed was served upon {defendant}.",
            f"{defendant} refused conveyance citing unverified force majeure grounds."
        ]
    else: # Commercial Disputes
        return [
            f"Formal procurement purchase order was placed by {defendant} with {plaintiff}.",
            f"Consignment was delivered, inspected, and accepted without recorded demur.",
            f"Tax invoices aggregating INR {claim_val:,.2f} were raised and accepted.",
            f"Pre-institution mediation under Section 12A of the Commercial Courts Act, 2015 was completed/exhausted.",
            f"Interest at statutory commercial rate is payable for default period."
        ]

def seed_cases():
    print(f"Loading {EXCEL_PATH}...")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    sheet = wb["Case_Dataset"]

    db = SessionLocal()

    # Fetch users
    judge = db.query(User).filter(User.role == RoleType.JUDGE).first()
    plaintiff_counsel = db.query(User).filter(User.email == "adv.himanshu@barcouncil.org").first()
    defendant_counsel = db.query(User).filter(User.email == "adv.verma@barcouncil.org").first()

    print("Reading rows and populating database...")
    seeded_count = 0

    for row_idx in range(2, sheet.max_row + 1):
        case_id_orig = sheet.cell(row=row_idx, column=1).value
        if not case_id_orig:
            continue

        category_str = sheet.cell(row=row_idx, column=2).value
        statute_str = sheet.cell(row=row_idx, column=3).value or ""
        court_str = sheet.cell(row=row_idx, column=4).value or "Commercial Court, Indore"
        filing_date_val = sheet.cell(row=row_idx, column=5).value
        plaintiff_str = clean_name(sheet.cell(row=row_idx, column=6).value)
        defendant_str = clean_name(sheet.cell(row=row_idx, column=8).value)
        claim_val_val = sheet.cell(row=row_idx, column=10).value or 0
        current_stage = sheet.cell(row=row_idx, column=11).value or "Plaint filed"
        cmh_date_val = sheet.cell(row=row_idx, column=12).value
        days_to_cmh = sheet.cell(row=row_idx, column=13).value
        issues_count = sheet.cell(row=row_idx, column=14).value or 2
        adjournments = sheet.cell(row=row_idx, column=15).value or 0
        adr_referred = sheet.cell(row=row_idx, column=16).value
        adr_mode = sheet.cell(row=row_idx, column=17).value
        hearing_mode = sheet.cell(row=row_idx, column=18).value or "Hybrid"
        priority_track = sheet.cell(row=row_idx, column=19).value or "Normal Track"
        case_status_raw = sheet.cell(row=row_idx, column=20).value or "Pending"

        # Date handling
        if isinstance(filing_date_val, datetime):
            filing_date = filing_date_val
        elif isinstance(filing_date_val, date):
            filing_date = datetime.combine(filing_date_val, time(10, 0))
        else:
            filing_date = datetime(2024, 1, 15, 10, 0)

        # 16-Digit CNR generation
        prefix = COURT_CNR_PREFIX.get(court_str, "MPIN01")
        seq_num = f"{seeded_count + 101:06d}"
        year_str = str(filing_date.year)
        cnr_number = f"{prefix}{seq_num}{year_str}"[:16]

        case_type = CATEGORY_MAP.get(category_str, CaseType.COMMERCIAL)
        case_title = f"{plaintiff_str} vs. {defendant_str}"
        claim_valuation = Decimal(str(claim_val_val))

        # Map CaseStatus
        if case_status_raw in ["Disposed", "Settled via ADR"]:
            status = CaseStatus.DISPOSED
        elif current_stage in ["Judgment reserved", "Arguments", "Issues framed", "Evidence stage"]:
            status = CaseStatus.CMO_ISSUED
        elif current_stage in ["Case Management Hearing"]:
            status = CaseStatus.CMH_SCHEDULED
        elif current_stage in ["Written statement filed", "Reply filed", "Accused appeared"]:
            status = CaseStatus.PRE_CMH
        else:
            status = CaseStatus.FILED

        # Check if already exists by CNR
        existing = db.query(Case).filter(Case.cnr_number == cnr_number).first()
        if existing:
            continue

        case_obj = Case(
            id=str(uuid.uuid4()),
            cnr_number=cnr_number,
            case_title=case_title,
            case_type=case_type,
            claim_valuation=claim_valuation,
            status=status,
            court_hall_id=court_str,
            judge_id=judge.id if judge else None,
            plaintiff_counsel_id=plaintiff_counsel.id if plaintiff_counsel else None,
            defendant_counsel_id=defendant_counsel.id if defendant_counsel else None,
            filing_date=filing_date
        )
        db.add(case_obj)
        db.flush()

        # 1. Pleadings
        plaint_hash = hashlib.sha256(f"{cnr_number}_PLAINT_{filing_date}".encode()).hexdigest()
        plaint_pleading = Pleading(
            id=str(uuid.uuid4()),
            case_id=case_obj.id,
            filed_by_id=plaintiff_counsel.id if plaintiff_counsel else (judge.id if judge else case_obj.id),
            pleading_type=PleadingType.PLAINT,
            file_storage_url=f"/storage/pleadings/{cnr_number}_plaint.pdf",
            sha256_checksum=plaint_hash,
            is_scrutinized=True,
            submitted_at=filing_date
        )
        db.add(plaint_pleading)
        db.flush()

        # Add exhibit
        exhibit_hash = hashlib.sha256(f"{cnr_number}_EXHIBIT_P1".encode()).hexdigest()
        exhibit = Exhibit(
            id=str(uuid.uuid4()),
            pleading_id=plaint_pleading.id,
            exhibit_number="P-1",
            document_title="Primary Agreement / Contract / Instrument",
            file_storage_url=f"/storage/exhibits/{cnr_number}_P1.pdf",
            sha256_checksum=exhibit_hash,
            uploaded_at=filing_date
        )
        db.add(exhibit)

        # 2. Admissions & Denials Checklist
        averments = get_averments_for_category(category_str, plaintiff_str, defendant_str, statute_str, claim_valuation)
        checklist = AdmissionsChecklist(
            id=str(uuid.uuid4()),
            case_id=case_obj.id,
            version_id=1,
            is_locked=True if status in [CaseStatus.CMO_ISSUED, CaseStatus.DISPOSED] else False,
            readiness_score=0.90 if status in [CaseStatus.CMO_ISSUED, CaseStatus.DISPOSED] else (0.50 if status == CaseStatus.CMH_SCHEDULED else 0.20),
            plaintiff_submitted_at=filing_date + timedelta(days=2),
            defendant_submitted_at=filing_date + timedelta(days=15) if status != CaseStatus.FILED else None
        )
        db.add(checklist)
        db.flush()

        # Add averment items
        for idx, stmt in enumerate(averments):
            if status in [CaseStatus.CMO_ISSUED, CaseStatus.DISPOSED]:
                if idx < 2:
                    p_stance = StanceType.ADMIT_CONTENTS
                    d_stance = StanceType.ADMIT_CONTENTS
                    res_status = ItemResolutionStatus.UNDISPUTED_FACT
                elif idx == 2:
                    p_stance = StanceType.ADMIT_CONTENTS
                    d_stance = StanceType.DENY
                    res_status = ItemResolutionStatus.CONTESTED_ISSUE
                else:
                    p_stance = StanceType.ADMIT_CONTENTS
                    d_stance = StanceType.DENY
                    res_status = ItemResolutionStatus.CONTESTED_ISSUE
            elif status == CaseStatus.CMH_SCHEDULED:
                p_stance = StanceType.ADMIT_CONTENTS
                d_stance = StanceType.ADMIT_CONTENTS if idx == 0 else StanceType.PENDING
                res_status = ItemResolutionStatus.UNDISPUTED_FACT if idx == 0 else ItemResolutionStatus.PENDING
            else:
                p_stance = StanceType.ADMIT_CONTENTS
                d_stance = StanceType.PENDING
                res_status = ItemResolutionStatus.PENDING

            item = AdmissionItem(
                id=str(uuid.uuid4()),
                checklist_id=checklist.id,
                statement_index=idx + 1,
                factual_statement=stmt,
                exhibit_ref_id=exhibit.id,
                plaintiff_stance=p_stance,
                defendant_stance=d_stance,
                resolution_status=res_status
            )
            db.add(item)

        # 3. CMH Session
        if cmh_date_val:
            if isinstance(cmh_date_val, datetime):
                cmh_date = cmh_date_val.date()
            elif isinstance(cmh_date_val, date):
                cmh_date = cmh_date_val
            else:
                cmh_date = filing_date.date() + timedelta(days=60)

            session_status = SessionStatus.CONCLUDED if status in [CaseStatus.CMO_ISSUED, CaseStatus.DISPOSED] else (SessionStatus.REFERRED_ADR if adr_referred == "Yes" else SessionStatus.SCHEDULED)
            cmh_session = CMHSession(
                id=str(uuid.uuid4()),
                case_id=case_obj.id,
                judge_id=judge.id if judge else None,
                court_hall_id=court_str,
                hearing_date=cmh_date,
                start_time=time(10, 30),
                end_time=time(11, 15),
                status=session_status,
                webrtc_room_id=f"webrtc-cmh-{cnr_number.lower()}"
            )
            db.add(cmh_session)

        # 4. CMO Order
        if status in [CaseStatus.CMO_ISSUED, CaseStatus.DISPOSED]:
            cmo_date = filing_date + timedelta(days=min(int(days_to_cmh or 90), 120))
            issues_list = [f"Issue #{i+1}: Whether the claim under {statute_str} is legally sustainable against {defendant_str}?" for i in range(min(int(issues_count or 2), 4))]
            cmo = CaseManagementOrder(
                id=str(uuid.uuid4()),
                case_id=case_obj.id,
                judge_id=judge.id if judge else None,
                order_date=cmo_date,
                undisputed_facts_summary=f"Parties concede commercial engagement, execution of primary instrument, and receipt of statutory notice under {statute_str}.",
                framed_issues_json=str(issues_list),
                witness_quota_plaintiff=2,
                witness_quota_defendant=2,
                pdf_storage_url=f"/storage/orders/CMO_{cnr_number}.pdf",
                sha256_digest=hashlib.sha256(f"CMO_{cnr_number}".encode()).hexdigest(),
                qr_verification_code=f"QR_SEAL_{cnr_number[:8]}",
                is_digitally_signed=True
            )
            db.add(cmo)
            db.flush()

            # Deadlines
            deadline1 = DiscoveryDeadline(
                id=str(uuid.uuid4()),
                order_id=cmo.id,
                task_description="Submission of Evidence Affidavits & Witness List",
                due_date=(cmo_date + timedelta(days=15)).date(),
                status=DeadlineStatus.COMPLETED if status == CaseStatus.DISPOSED else DeadlineStatus.PENDING
            )
            deadline2 = DiscoveryDeadline(
                id=str(uuid.uuid4()),
                order_id=cmo.id,
                task_description="Completion of Document Inspection & Cross-Examination",
                due_date=(cmo_date + timedelta(days=45)).date(),
                status=DeadlineStatus.COMPLETED if status == CaseStatus.DISPOSED else DeadlineStatus.PENDING
            )
            db.add(deadline1)
            db.add(deadline2)

        seeded_count += 1

    db.commit()
    db.close()
    print(f"🎉 Successfully seeded {seeded_count} cases from {EXCEL_PATH} into cmh_portal.db!")

if __name__ == "__main__":
    seed_cases()
