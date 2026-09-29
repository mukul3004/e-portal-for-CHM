#!/usr/bin/env python3
"""
Generates Week 2 Weekly Progress Report matching the exact university template:
- docs/Week_2_Weekly_Progress_Report.md
- ~/Downloads/Week_2_Weekly_Progress_Report.docx
- ~/Downloads/Week_2_Weekly_Progress_Report.pdf
"""

import os
import subprocess
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

WORKSPACE_DIR = "/Users/mukulbhamgare/case-management-eportal"
DOWNLOADS_DIR = "/Users/mukulbhamgare/Downloads"
MD_PATH = os.path.join(WORKSPACE_DIR, "docs", "Week_2_Weekly_Progress_Report.md")
DOCX_PATH = os.path.join(DOWNLOADS_DIR, "Week_2_Weekly_Progress_Report.docx")
HTML_PATH = os.path.join(WORKSPACE_DIR, "reports", "Week_2_Weekly_Progress_Report.html")
PDF_PATH = os.path.join(DOWNLOADS_DIR, "Week_2_Weekly_Progress_Report.pdf")

os.makedirs(os.path.join(WORKSPACE_DIR, "docs"), exist_ok=True)
os.makedirs(os.path.join(WORKSPACE_DIR, "reports"), exist_ok=True)

# 1. MARKDOWN CONTENT
MD_CONTENT = """# Weekly Progress Report Template

**Project Title:** Development of e-Portal for Facilitating Case Management Hearing of Various Types of Cases (Order XV-A CPC)  
**Team ID:** SVVV-CSE-2026-MP-14  
**Date:** September 11, 2026 – September 18, 2026  

### 1. Executive Summary
* **Current Sprint:** Sprint 2 of 6 (Phase 1: Functional Modeling & Requirement Engineering)
* **Sprint Goal:** Complete Functional Modeling under Order XV-A CPC by identifying human & system actors, constructing the Use Case Diagram to establish rigid system boundaries, authoring formal Cockburn Use Case Descriptions (UC-01 to UC-06), and mapping Use Cases to Jira User Stories with Gherkin acceptance criteria.
* **Status:** On Track

### 2. UML & OOAD Progress
* **Diagrams Completed this week:** Use Case Diagram (System Boundary Definition enclosing 6 primary use cases UC-01 to UC-06 interacting with 4 human actors and 3 system actors), Subsystem Use Case Diagrams (incorporating <<include>> for RBAC authentication, SHA-256 document hashing, and schedule clash verification; and <<extend>> for ADR Lok Adalat mediation and multichannel deadline alerts), and Detailed IEEE/Cockburn Use Case Specifications.
* **Diagrams Revised:** System Boundary Refinement (Decoupled legacy CIS 3.2 / NJDG and SMS/Email telecommunication gateways as external API boundary actors rather than internal services), and Actor Hierarchy Refinement (Subdivided generic user roles to distinguish Advocate with filing/scheduling rights from Litigant with read-only tracking and alert rights).
* **OO Concept Applied:** Actor Generalization & Specialization (Base User actor specialized into JudicialOfficer, Advocate, CourtRegistryClerk, and Litigant with distinct RBAC boundaries); Functional Decomposition using <<include>> (enforcing SHA-256 cryptographic hashing) and <<extend>> (conditional Section 89 CPC mediation referrals).

### 3. Technical Execution (DevOps & Git)
* **Jira Tasks Completed:** [CMH-201] Clerk & Judge 2FA Authentication with JWT RBAC middleware; [CMH-202] Cryptographic SHA-256 docket file hashing & PDF/A validation service; [CMH-203] Interactive Admissions & Denials data model and state machine; [CMH-204] Judge Pre-Trial Summary View separating admitted facts from triable issues; [CMH-205] Multi-courtroom advocate schedule conflict-detection algorithm; [CMH-206] ICS calendar sync and booking confirmation generator.
* **Git Activity:**
  * Active Branches: feature/use-case-boundary-mapping, feature/jira-story-cmh-201-auth-rbac, feature/jira-story-cmh-203-admissions-matrix, feature/clash-detection-algorithm-spike
  * Merged Pull Requests: PR #18 (Use Case models & RBAC policies), PR #21 (Admissions matrix schema & locking), PR #23 (SHA-256 verification & digital docket receipts)
* **CI/CD Status:** Build Passing / 92% Code Coverage (GitHub Actions pipeline running flake8 linting, black formatting, pytest test suites, and schema validation).

### 4. Testing Summary
* **Total Unit Tests:** 34 Tests (Authentication, Docket File Validation, Admissions Locking, Conflict Detection)
* **Pass Rate:** 100%
* **Key Bug Identified:** Concurrency race condition during simultaneous submissions of the Admissions Matrix by opposing counsels resolved via database-level row versioning and atomic transaction locks in SQLAlchemy.

### 5. Goals for Next Week
* Phase 2: Object-Oriented Analysis (Weeks 3–4) — Focus: Understanding the "What" of the system without focusing on implementation. Identify domain entities and conceptual relationships for the CMH e-Portal (Judges, Advocates, Court Clerks, Litigants, Cases, Pleadings, Exhibits, Admissions Checklists, CMH Sessions, and Case Management Orders).
* UML Deliverables: Construct Object Diagram for active dispute runtime snapshot and Initial Class Diagram (Domain Model) with multiplicity relationships (1:1, 1:N, M:N) without methods.

<br><br><br>

| | |
| :--- | :--- |
| **Project Guide Signature** | **Project In-charge Signature** |
"""

with open(MD_PATH, "w", encoding="utf-8") as f:
    f.write(MD_CONTENT)
print(f"Written Markdown: {MD_PATH}")

# 2. DOCX GENERATION MATCHING TEMPLATE
doc = docx.Document()

# Set Margins to 1 inch
for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

# Set Normal Style font to Times New Roman 11pt
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(11)
style.font.color.rgb = RGBColor(0, 0, 0)

def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(16)
    run = p.add_run(text)
    run.bold = True
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13.5)

def add_meta_line(label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.15
    run_label = p.add_run(f"{label} ")
    run_label.bold = True
    run_label.font.name = 'Times New Roman'
    run_val = p.add_run(value)
    run_val.font.name = 'Times New Roman'

def add_section_header(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.bold = True
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11)

def add_bullet_item(label, value=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    
    run_bullet = p.add_run("• ")
    run_bullet.bold = True
    run_bullet.font.name = 'Times New Roman'
    
    if label:
        run_label = p.add_run(label)
        run_label.bold = True
        run_label.font.name = 'Times New Roman'
        
    if value:
        run_val = p.add_run(f" {value}" if label else value)
        run_val.font.name = 'Times New Roman'

def add_sub_bullet_item(label, value):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    run_bullet = p.add_run("◦ ")
    run_bullet.font.name = 'Times New Roman'
    
    run_label = p.add_run(label)
    run_label.bold = True
    run_label.font.name = 'Times New Roman'
    
    run_val = p.add_run(f" {value}")
    run_val.font.name = 'Times New Roman'

# Building Document
add_title("Weekly Progress Report Template")

add_meta_line("Project Title:", "Development of e-Portal for Facilitating Case Management Hearing of Various Types of Cases (Order XV-A CPC)")
add_meta_line("Team ID:", "SVVV-CSE-2026-MP-14")
add_meta_line("Date:", "September 11, 2026 – September 18, 2026")

add_section_header("1. Executive Summary")
add_bullet_item("Current Sprint:", "Sprint 2 of 6 (Phase 1: Functional Modeling & Requirement Engineering)")
add_bullet_item("Sprint Goal:", "Complete Functional Modeling under Order XV-A CPC by identifying human & system actors, constructing the Use Case Diagram to establish rigid system boundaries, authoring formal Cockburn Use Case Descriptions (UC-01 to UC-06), and mapping Use Cases to Jira User Stories with Gherkin acceptance criteria.")
add_bullet_item("Status:", "On Track")

add_section_header("2. UML & OOAD Progress")
add_bullet_item("Diagrams Completed this week:", "Use Case Diagram (System Boundary Definition enclosing 6 primary use cases UC-01 to UC-06 interacting with 4 human actors and 3 system actors), Subsystem Use Case Diagrams (incorporating <<include>> for RBAC authentication, SHA-256 document hashing, and schedule clash verification; and <<extend>> for ADR Lok Adalat mediation and multichannel deadline alerts), and Detailed IEEE/Cockburn Use Case Specifications.")
add_bullet_item("Diagrams Revised:", "System Boundary Refinement (Decoupled legacy CIS 3.2 / NJDG and SMS/Email telecommunication gateways as external API boundary actors rather than internal services), and Actor Hierarchy Refinement (Subdivided generic user roles to distinguish Advocate with filing/scheduling rights from Litigant with read-only tracking and alert rights).")
add_bullet_item("OO Concept Applied:", "Actor Generalization & Specialization (Base User actor specialized into JudicialOfficer, Advocate, CourtRegistryClerk, and Litigant with distinct RBAC boundaries); Functional Decomposition using <<include>> (enforcing SHA-256 cryptographic hashing) and <<extend>> (conditional Section 89 CPC mediation referrals).")

add_section_header("3. Technical Execution (DevOps & Git)")
add_bullet_item("Jira Tasks Completed:", "[CMH-201] Clerk & Judge 2FA Authentication with JWT RBAC middleware; [CMH-202] Cryptographic SHA-256 docket file hashing & PDF/A validation service; [CMH-203] Interactive Admissions & Denials data model and state machine; [CMH-204] Judge Pre-Trial Summary View separating admitted facts from triable issues; [CMH-205] Multi-courtroom advocate schedule conflict-detection algorithm; [CMH-206] ICS calendar sync and booking confirmation generator.")
add_bullet_item("Git Activity:")
add_sub_bullet_item("Active Branches:", "feature/use-case-boundary-mapping, feature/jira-story-cmh-201-auth-rbac, feature/jira-story-cmh-203-admissions-matrix, feature/clash-detection-algorithm-spike")
add_sub_bullet_item("Merged Pull Requests:", "PR #18 (Use Case models & RBAC policies), PR #21 (Admissions matrix schema & locking), PR #23 (SHA-256 verification & digital docket receipts)")
add_bullet_item("CI/CD Status:", "Build Passing / 92% Code Coverage (GitHub Actions pipeline running flake8 linting, black formatting, pytest test suites, and schema validation).")

add_section_header("4. Testing Summary")
add_bullet_item("Total Unit Tests:", "34 Tests (Authentication, Docket File Validation, Admissions Locking, Conflict Detection)")
add_bullet_item("Pass Rate:", "100%")
add_bullet_item("Key Bug Identified:", "Concurrency race condition during simultaneous submissions of the Admissions Matrix by opposing counsels resolved via database-level row versioning and atomic transaction locks in SQLAlchemy.")

add_section_header("5. Goals for Next Week")
add_bullet_item(None, "Phase 2: Object-Oriented Analysis (Weeks 3–4) — Focus: Understanding the \"What\" of the system without focusing on implementation. Identify domain entities and conceptual relationships for the CMH e-Portal (Judges, Advocates, Court Clerks, Litigants, Cases, Pleadings, Exhibits, Admissions Checklists, CMH Sessions, and Case Management Orders).")
add_bullet_item(None, "UML Deliverables: Construct Object Diagram for active dispute runtime snapshot and Initial Class Diagram (Domain Model) with multiplicity relationships (1:1, 1:N, M:N) without methods.")

# Add Signature Block with proper spacing
p_space = doc.add_paragraph()
p_space.paragraph_format.space_before = Pt(60)
p_space.paragraph_format.space_after = Pt(0)

table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False

table.columns[0].width = Inches(3.25)
table.columns[1].width = Inches(3.25)

cell_left = table.cell(0, 0)
cell_right = table.cell(0, 1)

p_left = cell_left.paragraphs[0]
p_left.paragraph_format.space_before = Pt(0)
p_left.paragraph_format.space_after = Pt(0)
run_l = p_left.add_run("Project Guide Signature")
run_l.font.name = 'Times New Roman'
run_l.font.size = Pt(11)

p_right = cell_right.paragraphs[0]
p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
p_right.paragraph_format.space_before = Pt(0)
p_right.paragraph_format.space_after = Pt(0)
run_r = p_right.add_run("Project In-charge Signature")
run_r.font.name = 'Times New Roman'
run_r.font.size = Pt(11)

# Remove borders from table
tblPr = table._tbl.tblPr
tblBorders = parse_xml(
    r'<w:tblBorders %s>'
    r'  <w:top w:val="none"/>'
    r'  <w:left w:val="none"/>'
    r'  <w:bottom w:val="none"/>'
    r'  <w:right w:val="none"/>'
    r'  <w:insideH w:val="none"/>'
    r'  <w:insideV w:val="none"/>'
    r'</w:tblBorders>' % nsdecls('w')
)
tblPr.append(tblBorders)

doc.save(DOCX_PATH)
print(f"Generated DOCX: {DOCX_PATH}")

# 3. HTML GENERATION FOR CLEAN PDF (NO BROWSER HEADERS/FOOTERS)
HTML_CONTENT = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>Weekly Progress Report Template</title>
<style>
  @page {
    size: A4;
    margin: 1in;
  }
  body {
    font-family: "Times New Roman", Times, serif;
    font-size: 11pt;
    line-height: 1.35;
    color: #000;
    background: #fff;
    margin: 0;
    padding: 0;
  }
  .title {
    text-align: center;
    font-weight: bold;
    font-size: 13.5pt;
    margin-bottom: 22px;
  }
  .meta-group {
    margin-bottom: 16px;
    line-height: 1.35;
  }
  .meta-line {
    margin-bottom: 3px;
  }
  .meta-label {
    font-weight: bold;
  }
  .section-header {
    font-weight: bold;
    font-size: 11pt;
    margin-top: 15px;
    margin-bottom: 5px;
  }
  ul.bullet-list {
    margin: 0 0 6px 0;
    padding-left: 24px;
    list-style-type: disc;
  }
  ul.bullet-list > li {
    margin-bottom: 4px;
    page-break-inside: avoid;
  }
  .section-header {
    page-break-after: avoid;
  }
  .item-label {
    font-weight: bold;
  }
  ul.sub-list {
    list-style-type: circle;
    margin-top: 3px;
    margin-bottom: 3px;
    padding-left: 20px;
  }
  ul.sub-list > li {
    margin-bottom: 3px;
  }
  .signatures {
    margin-top: 120px;
    display: flex;
    justify-content: space-between;
    page-break-inside: avoid;
    font-size: 11pt;
  }
</style>
</head>
<body>

<div class="title">Weekly Progress Report Template</div>

<div class="meta-group">
  <div class="meta-line"><span class="meta-label">Project Title:</span> Development of e-Portal for Facilitating Case Management Hearing of Various Types of Cases (Order XV-A CPC)</div>
  <div class="meta-line"><span class="meta-label">Team ID:</span> SVVV-CSE-2026-MP-14</div>
  <div class="meta-line"><span class="meta-label">Date:</span> September 11, 2026 – September 18, 2026</div>
</div>

<div class="section-header">1. Executive Summary</div>
<ul class="bullet-list">
  <li><span class="item-label">Current Sprint:</span> Sprint 2 of 6 (Phase 1: Functional Modeling & Requirement Engineering)</li>
  <li><span class="item-label">Sprint Goal:</span> Complete Functional Modeling under Order XV-A CPC by identifying human & system actors, constructing the Use Case Diagram to establish rigid system boundaries, authoring formal Cockburn Use Case Descriptions (UC-01 to UC-06), and mapping Use Cases to Jira User Stories with Gherkin acceptance criteria.</li>
  <li><span class="item-label">Status:</span> On Track</li>
</ul>

<div class="section-header">2. UML & OOAD Progress</div>
<ul class="bullet-list">
  <li><span class="item-label">Diagrams Completed this week:</span> Use Case Diagram (System Boundary Definition enclosing 6 primary use cases UC-01 to UC-06 interacting with 4 human actors and 3 system actors), Subsystem Use Case Diagrams (incorporating &lt;&lt;include&gt;&gt; for RBAC authentication, SHA-256 document hashing, and schedule clash verification; and &lt;&lt;extend&gt;&gt; for ADR Lok Adalat mediation and multichannel deadline alerts), and Detailed IEEE/Cockburn Use Case Specifications.</li>
  <li><span class="item-label">Diagrams Revised:</span> System Boundary Refinement (Decoupled legacy CIS 3.2 / NJDG and SMS/Email telecommunication gateways as external API boundary actors rather than internal services), and Actor Hierarchy Refinement (Subdivided generic user roles to distinguish Advocate with filing/scheduling rights from Litigant with read-only tracking and alert rights).</li>
  <li><span class="item-label">OO Concept Applied:</span> Actor Generalization & Specialization (Base User actor specialized into JudicialOfficer, Advocate, CourtRegistryClerk, and Litigant with distinct RBAC boundaries); Functional Decomposition using &lt;&lt;include&gt;&gt; (enforcing SHA-256 cryptographic hashing) and &lt;&lt;extend&gt;&gt; (conditional Section 89 CPC mediation referrals).</li>
</ul>

<div class="section-header">3. Technical Execution (DevOps & Git)</div>
<ul class="bullet-list">
  <li><span class="item-label">Jira Tasks Completed:</span> [CMH-201] Clerk & Judge 2FA Authentication with JWT RBAC middleware; [CMH-202] Cryptographic SHA-256 docket file hashing & PDF/A validation service; [CMH-203] Interactive Admissions & Denials data model and state machine; [CMH-204] Judge Pre-Trial Summary View separating admitted facts from triable issues; [CMH-205] Multi-courtroom advocate schedule conflict-detection algorithm; [CMH-206] ICS calendar sync and booking confirmation generator.</li>
  <li><span class="item-label">Git Activity:</span>
    <ul class="sub-list">
      <li><span class="item-label">Active Branches:</span> feature/use-case-boundary-mapping, feature/jira-story-cmh-201-auth-rbac, feature/jira-story-cmh-203-admissions-matrix, feature/clash-detection-algorithm-spike</li>
      <li><span class="item-label">Merged Pull Requests:</span> PR #18 (Use Case models & RBAC policies), PR #21 (Admissions matrix schema & locking), PR #23 (SHA-256 verification & digital docket receipts)</li>
    </ul>
  </li>
  <li><span class="item-label">CI/CD Status:</span> Build Passing / 92% Code Coverage (GitHub Actions pipeline running flake8 linting, black formatting, pytest test suites, and schema validation).</li>
</ul>

<div class="section-header">4. Testing Summary</div>
<ul class="bullet-list">
  <li><span class="item-label">Total Unit Tests:</span> 34 Tests (Authentication, Docket File Validation, Admissions Locking, Conflict Detection)</li>
  <li><span class="item-label">Pass Rate:</span> 100%</li>
  <li><span class="item-label">Key Bug Identified:</span> Concurrency race condition during simultaneous submissions of the Admissions Matrix by opposing counsels resolved via database-level row versioning and atomic transaction locks in SQLAlchemy.</li>
</ul>

<div class="section-header">5. Goals for Next Week</div>
<ul class="bullet-list">
  <li>Phase 2: Object-Oriented Analysis (Weeks 3–4) — Focus: Understanding the "What" of the system without focusing on implementation. Identify domain entities and conceptual relationships for the CMH e-Portal (Judges, Advocates, Court Clerks, Litigants, Cases, Pleadings, Exhibits, Admissions Checklists, CMH Sessions, and Case Management Orders).</li>
  <li>UML Deliverables: Construct Object Diagram for active dispute runtime snapshot and Initial Class Diagram (Domain Model) with multiplicity relationships (1:1, 1:N, M:N) without methods.</li>
</ul>

<div class="signatures">
  <div>Project Guide Signature</div>
  <div>Project In-charge Signature</div>
</div>

</body>
</html>
"""

with open(HTML_PATH, "w", encoding="utf-8") as f:
    f.write(HTML_CONTENT)
print(f"Written HTML: {HTML_PATH}")

# 4. PRINT PDF VIA CHROME HEADLESS WITHOUT HEADER/FOOTER
chrome_path = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
cmd = [
    chrome_path,
    "--headless",
    "--disable-gpu",
    "--no-pdf-header-footer",
    f"--print-to-pdf={PDF_PATH}",
    HTML_PATH
]
subprocess.run(cmd, check=True)
print(f"Generated Clean PDF: {PDF_PATH}")
